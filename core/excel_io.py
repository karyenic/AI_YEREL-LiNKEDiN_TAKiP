"""Excel içe/dışa aktarma modülü."""
import io
import hashlib
import unicodedata
from datetime import datetime
from pathlib import Path
import openpyxl
import pandas as pd
from core.database import aday_ekle, adaylari_getir


# ═══════════════════════════════════════════════════════════
# YARDIMCI FONKSİYONLAR
# ═══════════════════════════════════════════════════════════
def _satir_hash(isim, tarih, aciklama):
    """Mükerrer kontrolü için hash. Tarih DAHİL EDİLMEZ."""
    def normalize(s):
        if not s:
            return ""
        s = str(s)
        s = unicodedata.normalize("NFKD", s)
        # Görünmez karakterleri temizle
        for ch in ["\u00a0", "\u200b", "\u200c", "\u200d", "\ufeff", "\u2028", "\u2029"]:
            s = s.replace(ch, " ")
        # Tüm whitespace'i tek boşluğa indir + baş/son temizle
        s = " ".join(s.split())
        return s.lower()

    isim_norm = normalize(isim)
    aciklama_norm = normalize(aciklama)
    raw = f"{isim_norm}|{aciklama_norm}".encode("utf-8")
    return hashlib.md5(raw).hexdigest()


def _kolon_bul(headers, *names):
    """Sütun başlığı arar. Tam eşleşme, sonra kısmi eşleşme."""
    for n in names:
        n = n.lower().strip()
        # 1. Tam eşleşme
        if n in headers:
            return headers.index(n)
        # 2. Kısmi eşleşme (başlık, aranan kelimeyi içeriyor mu?)
        for i, h in enumerate(headers):
            if n in h:
                return i
    return None


def _evet_mi(val):
    """Excel hücresini 0/1'e çevirir."""
    if val is None:
        return 0
    s = str(val).strip().lower()
    return 1 if s in ["evet", "yes", "true", "1", "var", "x", "✓", "doğru", "e"] else 0


# ═══════════════════════════════════════════════════════════
# EXCEL İÇE AKTARMA
# ═══════════════════════════════════════════════════════════
def excel_ice_aktar(excel_yolu):
    """Excel dosyasını okur ve yeni adayları DB'ye ekler."""
    yol = Path(excel_yolu)
    if not yol.exists():
        return False, f"Excel dosyası bulunamadı: {yol}"

    try:
        wb = openpyxl.load_workbook(yol, data_only=True)
        ws = wb.active

        # 1. Başlıkları oku (1. satır)
        headers = []
        for c in ws[1]:
            if c.value is None:
                headers.append("")
            else:
                headers.append(str(c.value).strip().lower())

        # 2. Sütun indekslerini bul
        isim_i = _kolon_bul(headers, "isim", "ad", "adi", "adi soyadi", "name", "aday", "aday ismi")
        aciklama_i = _kolon_bul(headers, "aciklama", "dusunceler", "not", "notes", "thoughts", "detay")
        linkedin_i = _kolon_bul(headers, "linkedin url", "linkedin", "url", "link", "profil")
        if isim_i is None:
            return False, f"Excel'de 'İsim' sütunu bulunamadı. Başlık satırı: {headers}"

        # 3. Satırları işle
        eklenen = 0
        atlanan = 0
        bugun = datetime.now().strftime("%d %m %y")

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue

            isim = str(row[isim_i]).strip() if isim_i < len(row) and row[isim_i] else ""
            if not isim or isim.isdigit():
                continue

            # Tarih
            # Tarih: HER ZAMAN sistem tarihi (Excel tarihi yok sayilir)

            tarih = bugun

            # Açıklama
            aciklama = ""
            if aciklama_i is not None and aciklama_i < len(row) and row[aciklama_i]:
                aciklama = str(row[aciklama_i]).strip()

            linkedin_url = None
            if linkedin_i is not None and linkedin_i < len(row) and row[linkedin_i]:
                linkedin_url = str(row[linkedin_i]).strip() or None

            h = _satir_hash(isim, tarih, aciklama)
            ok = aday_ekle(isim, tarih, aciklama, 0, 0, 0,
                           0, 0, 0, 0, kaynak_hash=h,
                           linkedin_url=linkedin_url)
            if ok:
                eklenen += 1
            else:
                atlanan += 1

        return True, f"{eklenen} yeni aday eklendi, {atlanan} zaten mevcuttu."

    except Exception as e:
        return False, f"Excel okuma hatası: {e}"


# ═══════════════════════════════════════════════════════════
# EXCEL DIŞA AKTARMA
# ═══════════════════════════════════════════════════════════
def excel_disa_aktar():
    """Adayları Excel formatında bytes olarak döner."""
    adaylar = adaylari_getir()
    if not adaylar:
        return None

    df = pd.DataFrame(adaylar)

    mapping = {
        'isim': 'ADI SOYADI',
        'linkedin_url': 'LINKEDIN URL',
        'aciklama': 'ACIKLAMA',
    }
    df = df.rename(columns=mapping)

    keep = [c for c in [
        'ADI SOYADI', 'LINKEDIN URL', 'ACIKLAMA'
    ] if c in df.columns]
    df = df[keep]

    out = io.BytesIO()
    with pd.ExcelWriter(out, engine='xlsxwriter') as writer:
        df.to_excel(writer, index=False, sheet_name='Adaylar')
    out.seek(0)
    return out.getvalue()