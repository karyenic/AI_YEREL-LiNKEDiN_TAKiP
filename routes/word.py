"""Word (DOCX) dışa aktarma."""
from flask import Blueprint, send_file, jsonify
import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from core.database import adaylari_getir

bp = Blueprint("word", __name__, url_prefix="/api/word")


@bp.route("/export", methods=["GET"])
def export():
    adaylar = adaylari_getir()
    if not adaylar:
        return jsonify({"error": "Veri yok"}), 404

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    baslik = doc.add_heading("LinkedIn Takip - Aday Listesi", level=1)
    baslik.alignment = WD_ALIGN_PARAGRAPH.CENTER

    tarih_p = doc.add_paragraph()
    tarih_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tarih_run = tarih_p.add_run(
        f"Yazdırma tarihi: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    )
    tarih_run.font.size = Pt(9)
    tarih_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph()

    sutunlar = ["ID", "İsim", "Tarih", "Açıklama",
                "Davet", "Randevu", "Plan", "Kayıt",
                "Takip", "Hayır", "İş Arıyor"]

    tablo = doc.add_table(rows=1, cols=len(sutunlar))
    tablo.style = "Light Grid Accent 1"
    tablo.alignment = WD_TABLE_ALIGNMENT.CENTER

    baslik_satiri = tablo.rows[0].cells
    for i, sutun_adi in enumerate(sutunlar):
        baslik_satiri[i].text = sutun_adi
        for p in baslik_satiri[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)

    for a in adaylar:
        satir = tablo.add_row().cells
        satir[0].text = str(a.get("id", ""))
        satir[1].text = str(a.get("isim", ""))
        satir[2].text = str(a.get("tarih", ""))
        satir[3].text = str(a.get("aciklama", ""))[:80]
        satir[4].text = "✓" if a.get("davet") == 1 else "✗"
        satir[5].text = "✓" if a.get("randevu") == 1 else "✗"
        satir[6].text = "✓" if a.get("plan") == 1 else "✗"
        satir[7].text = "✓" if a.get("kayit") == 1 else "✗"
        satir[8].text = "✓" if a.get("takip") == 1 else "✗"
        satir[9].text = "✓" if a.get("hayir") == 1 else "✗"
        satir[10].text = "✓" if a.get("is_ariyor") == 1 else "✗"
        for h in satir:
            for p in h.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)

    doc.add_paragraph()
    toplam_p = doc.add_paragraph()
    tr = toplam_p.add_run(f"Toplam: {len(adaylar)} aday")
    tr.font.bold = True
    tr.font.size = Pt(10)

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)

    return send_file(
        out,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=f"linkedin_adaylar_{datetime.now().strftime('%Y%m%d')}.docx"
    )
